"""Render a requirements .md to .docx in the house style used by docs/*.docx.

House style (matched from the existing files): Heading 1/2/3, Normal body,
'Light Grid Accent 1' tables with a bold header row, monospace code blocks.
Inline **bold**, `code` and *italic* are rendered as runs rather than left as
literal markers (the previous generator leaked them into the text).
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*|~~.+?~~)', re.S)

def add_runs(par, text, bold=False, italic=False, strike=False):
    """Split markdown inline markup into formatted runs (nesting-aware)."""
    def plain(t):
        if not t:
            return
        r = par.add_run(t)
        r.bold, r.italic, r.font.strike = bold or None, italic or None, strike or None

    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            add_runs(par, tok[2:-2], True, italic, strike)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.font.name = 'Consolas'; r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
            r.bold, r.italic = bold or None, italic or None
        elif tok.startswith('~~') and tok.endswith('~~') and len(tok) > 4:
            add_runs(par, tok[2:-2], bold, italic, True)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            add_runs(par, tok[1:-1], bold, True, strike)
        else:
            plain(tok)

def split_row(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    return [c.strip() for c in line.split('|')]

def is_sep(line):
    return bool(re.fullmatch(r'\|?[\s:|-]+\|[\s:|-]*', line.strip())) and '-' in line

def convert(md_path, docx_path):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()
    i, n = 0, len(lines)
    para_buf = []

    def flush():
        nonlocal para_buf
        if para_buf:
            p = doc.add_paragraph()
            add_runs(p, ' '.join(para_buf))
            para_buf = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:                                   # blank -> end paragraph
            flush(); i += 1; continue

        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)       # heading
        if m:
            flush()
            lvl = min(len(m.group(1)), 3)
            h = doc.add_heading(level=lvl)
            add_runs(h, m.group(2))
            i += 1; continue

        if stripped.startswith('```'):                     # fenced code
            flush(); i += 1
            code = []
            while i < n and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i += 1
            i += 1
            for cl in code:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(cl if cl.strip() else ' ')
                r.font.name = 'Consolas'; r.font.size = Pt(9)
            doc.add_paragraph()
            continue

        if stripped.startswith('|') and i + 1 < n and is_sep(lines[i + 1]):   # table
            flush()
            header = split_row(lines[i]); i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i])); i += 1
            cols = max([len(header)] + [len(r) for r in rows])
            # A metadata table whose header row is entirely empty (| | |) is
            # rendered without that row, matching the original generator.
            has_header = any(c.strip() for c in header)
            t = doc.add_table(rows=1 if has_header else 0, cols=cols)
            try: t.style = 'Light Grid Accent 1'
            except KeyError: t.style = 'Table Grid'
            if has_header:
                for c, txt in enumerate(header + [''] * (cols - len(header))):
                    cell = t.rows[0].cells[c]; cell.text = ''
                    add_runs(cell.paragraphs[0], txt)
                    for r in cell.paragraphs[0].runs: r.bold = True
            for row in rows:
                cells = t.add_row().cells
                for c, txt in enumerate(row + [''] * (cols - len(row))):
                    cells[c].text = ''
                    add_runs(cells[c].paragraphs[0], txt)
            doc.add_paragraph()
            continue

        m_b = re.match(r'^(\s*)[-*+]\s+(.*)$', line)        # bullet
        m_n = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)      # numbered
        if m_b or m_n:
            flush()
            indent = len((m_b or m_n).group(1))
            text = m_b.group(2) if m_b else m_n.group(3)
            i += 1
            # absorb wrapped continuation lines (indented, not a new item,
            # not a table/heading/fence/blank) so the item stays one paragraph
            while i < n:
                nxt = lines[i]
                if (not nxt.strip() or re.match(r'^\s*([-*+]|\d+\.)\s', nxt)
                        or nxt.lstrip().startswith(('|', '#', '```', '>'))
                        or (len(nxt) - len(nxt.lstrip())) <= indent):
                    break
                text += ' ' + nxt.strip(); i += 1
            if m_b:
                depth = min(indent // 2, 2)
                style = 'List Bullet' if depth == 0 else f'List Bullet {depth+1}'
            else:
                style = 'List Number'
            try: p = doc.add_paragraph(style=style)
            except KeyError: p = doc.add_paragraph(style='List Bullet')
            add_runs(p, text); continue

        if stripped.startswith('>'):                        # blockquote (may wrap)
            flush()
            quote = []
            while i < n and lines[i].strip().startswith('>'):
                quote.append(lines[i].strip().lstrip('>').strip()); i += 1
            try: p = doc.add_paragraph(style='Quote')
            except KeyError: p = doc.add_paragraph()
            add_runs(p, ' '.join(q for q in quote if q)); continue

        if re.fullmatch(r'-{3,}|_{3,}|\*{3,}', stripped):   # hr
            flush(); doc.add_paragraph(); i += 1; continue

        para_buf.append(stripped); i += 1

    flush()
    doc.save(docx_path)

if __name__ == '__main__':
    convert(sys.argv[1], sys.argv[2])
    print('wrote', sys.argv[2])
